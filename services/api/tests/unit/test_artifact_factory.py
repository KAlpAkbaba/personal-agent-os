"""Unit tests: app.artifacts.factory (docs/M22_ARTIFACT_FACTORY_SPEC.md, ADR-0085).

create() -> renders -> validations, against SQLite + InMemoryObjectStore (the same
harness tests/unit/test_artifact_service.py already uses for the M13 render-store
bridge). Idempotency on the spec's own content_hash; an invalid render is stored
"invalid" and named, never silently upgraded to "done".
"""

from __future__ import annotations

import json

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

from app.artifacts import factory, renderers, service
from app.artifacts.models import (
    ARTIFACT_STATE_READY,
    CANONICAL_FORMAT_ARTIFACT_SPEC_JSON,
    RENDER_STATE_INVALID,
    RENDER_STATE_VALID,
    Artifact,
    ArtifactRender,
    ArtifactVersion,
    ResearchSource,
    Task,
    TaskRun,
)
from app.artifacts.spec import ArtifactSpec
from app.object_store import InMemoryObjectStore

ARTIFACT_TABLES = [
    Task.__table__,
    TaskRun.__table__,
    Artifact.__table__,
    ArtifactVersion.__table__,
    ArtifactRender.__table__,
    ResearchSource.__table__,
]


@pytest.fixture()
def db() -> Session:
    engine = create_engine("sqlite://")
    for table in ARTIFACT_TABLES:
        table.create(engine)
    session_factory = sessionmaker(bind=engine, expire_on_commit=False)
    session = session_factory()
    yield session
    session.close()
    engine.dispose()


@pytest.fixture()
def store() -> InMemoryObjectStore:
    return InMemoryObjectStore()


def _spec(path: str) -> ArtifactSpec:
    with open(path, encoding="utf-8") as f:
        return ArtifactSpec.model_validate(json.load(f))


BUDGET = "tests/fixtures/artifacts/specs/butce-tablosu.json"
PRESENTATION = "tests/fixtures/artifacts/specs/q3-sunum.json"
PAGE = "tests/fixtures/artifacts/specs/hosgeldin-sayfasi.json"


# --------------------------------------------------------------------- create()


def test_create_renders_every_format_the_kind_produces(db: Session, store) -> None:
    spec = _spec(BUDGET)
    result = factory.create(db, store, spec=spec)
    assert result.created is True
    assert {r.format for r in result.renders} == set(spec.formats())
    assert result.all_valid
    assert result.kind == "spreadsheet"
    assert result.title == "Bütçe 2026"


def test_create_persists_canonical_json_not_markdown(db: Session, store) -> None:
    spec = _spec(PAGE)
    result = factory.create(db, store, spec=spec)
    artifact = service.get_artifact(db, result.artifact_id)
    assert artifact.canonical_format == CANONICAL_FORMAT_ARTIFACT_SPEC_JSON
    version = service.get_current_version(db, artifact.id)
    reparsed = ArtifactSpec.model_validate_json(version.canonical_body)
    assert reparsed.title == spec.title
    assert reparsed.kind == spec.kind


def test_create_reaches_artifact_ready_state(db: Session, store) -> None:
    spec = _spec(PAGE)
    result = factory.create(db, store, spec=spec)
    artifact = service.get_artifact(db, result.artifact_id)
    assert artifact.state == ARTIFACT_STATE_READY


def test_create_stores_bytes_reopenable_from_the_store(db: Session, store) -> None:
    spec = _spec(PRESENTATION)
    result = factory.create(db, store, spec=spec)
    version = service.get_current_version(db, result.artifact_id)
    row = service.get_render(db, version.id, "pptx")
    data = store.get(row.object_key)
    assert data[:2] == b"PK"
    assert len(data) == row.size_bytes


def test_every_render_row_carries_its_validation_report(db: Session, store) -> None:
    spec = _spec(BUDGET)
    result = factory.create(db, store, spec=spec)
    version = service.get_current_version(db, result.artifact_id)
    for fmt in spec.formats():
        row = service.get_render(db, version.id, fmt)
        assert row.state == RENDER_STATE_VALID
        assert row.validation_json is not None
        assert row.validation_json["ok"] is True


# ------------------------------------------------------------------- idempotency


def test_create_is_idempotent_on_the_same_spec(db: Session, store) -> None:
    spec = _spec(BUDGET)
    first = factory.create(db, store, spec=spec)
    second = factory.create(db, store, spec=spec)
    assert second.created is False
    assert second.artifact_id == first.artifact_id
    assert second.version == first.version
    assert len(service.list_artifacts(db)) == 1


def test_create_with_a_different_spec_makes_a_new_artifact(db: Session, store) -> None:
    a = factory.create(db, store, spec=_spec(BUDGET))
    b = factory.create(db, store, spec=_spec(PRESENTATION))
    assert a.artifact_id != b.artifact_id
    assert len(service.list_artifacts(db)) == 2


# --------------------------------------------------------------- invalid renders


class _LyingXlsxRenderer:
    """A renderer that silently drops the last data row -- stands in for a real
    corruption to prove the factory records an INVALID render rather than treating
    every render as automatically correct."""

    format = "xlsx"
    mime_type = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    def render_spec(self, spec):
        from openpyxl import Workbook

        wb = Workbook()
        wb.remove(wb.active)
        ws = wb.create_sheet("Ozet")
        ws.cell(row=1, column=1, value="Kalem")
        ws.cell(row=1, column=2, value="Tutar")
        # Only the first data row -- the rest are dropped.
        ws.cell(row=2, column=1, value=spec.sheets[0].rows[0][0])
        ws.cell(row=2, column=2, value=spec.sheets[0].rows[0][1])
        import io

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()


def test_a_lying_renderer_is_recorded_invalid_with_failing_refs(
    db: Session, store, monkeypatch
) -> None:
    monkeypatch.setitem(renderers._FACTORY_RENDERERS, "xlsx", _LyingXlsxRenderer())
    spec = _spec(BUDGET)
    result = factory.create(db, store, spec=spec)
    xlsx = next(r for r in result.renders if r.format == "xlsx")
    assert xlsx.state == RENDER_STATE_INVALID
    assert xlsx.failing_refs  # at least one ref named
    assert any("A3" in ref or "B3" in ref for ref in xlsx.failing_refs)
    assert result.all_valid is False
    assert result.failing == [xlsx]
    # the csv render (untouched) is still valid -- one bad format never poisons
    # the whole artifact.
    csv_render = next(r for r in result.renders if r.format == "csv")
    assert csv_render.valid


# ------------------------------------------------------- validation lookup / revalidate


def test_render_validation_returns_the_stored_report(db: Session, store) -> None:
    spec = _spec(BUDGET)
    result = factory.create(db, store, spec=spec)
    payload = factory.render_validation(db, artifact_id=result.artifact_id, fmt="xlsx")
    assert payload["state"] == "valid"
    assert payload["validation"]["ok"] is True
    assert payload["validation"]["checks"]


def test_render_validation_is_none_for_unknown_artifact(db: Session) -> None:
    import uuid

    assert factory.render_validation(db, artifact_id=uuid.uuid4(), fmt="xlsx") is None


def test_revalidate_forces_a_fresh_render_and_validate_pass(db: Session, store) -> None:
    spec = _spec(BUDGET)
    result = factory.create(db, store, spec=spec)
    version = service.get_current_version(db, result.artifact_id)
    row = service.get_render(db, version.id, "xlsx")
    store.delete(row.object_key)  # simulate the object having gone missing
    payload = factory.revalidate(db, store, artifact_id=result.artifact_id, fmt="xlsx")
    assert payload["state"] == "valid"
    assert store.exists(row.object_key)


def test_download_path_matches_the_existing_m13_render_route() -> None:
    import uuid

    artifact_id = uuid.uuid4()
    assert factory.download_path(artifact_id, "xlsx") == f"/v1/artifacts/{artifact_id}/renders/xlsx"
