"""Unit tests: independent-reader validation (docs/M22_ARTIFACT_FACTORY_SPEC.md §3,
ADR-0085 decision 3).

Three properties, per the task brief:

1. Every fixture render passes its own ``truth.json.expect`` (the oracle a
   different script built independently of this module).
2. An injected "lying renderer" — one that drops a row / a slide / a heading — is
   caught by validate(), with the correct failing ref named.
3. Every third-party format parsed here (DOCX/XLSX/PPTX/PDF) has a resource bound,
   proven against a zip-bomb-shaped OOXML package, a PDF-shaped bound and an
   oversized plain-text/JSON input — never an unhandled exception.
"""

from __future__ import annotations

import glob
import io
import json
import zipfile

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from app.artifacts import renderers as R
from app.artifacts import validation as V
from app.artifacts.spec import ArtifactSpec, Sheet

FIXTURE_SPECS = sorted(glob.glob("tests/fixtures/artifacts/specs/*.json"))
TRUTH = json.load(open("tests/fixtures/artifacts/truth.json", encoding="utf-8"))["specs"]


def _load_spec(path: str) -> ArtifactSpec:
    with open(path, encoding="utf-8") as f:
        return ArtifactSpec.model_validate(json.load(f))


def _fixture_name(path: str) -> str:
    return path.replace("\\", "/").rsplit("/", 1)[-1][: -len(".json")]


def _cases() -> list[tuple[str, str]]:
    out = []
    for path in FIXTURE_SPECS:
        spec = _load_spec(path)
        for fmt in spec.formats():
            out.append((path, fmt))
    return out


CASES = _cases()


def _matching_checks(report: V.ValidationReport, ref: str | None) -> list[dict]:
    if ref is None:
        return report.checks
    return [c for c in report.checks if c.get("ref") == ref]


def _truth_entry_satisfied(report: V.ValidationReport, entry: dict) -> bool:
    ref = entry.get("ref")
    matched = _matching_checks(report, ref)
    if "contains" in entry:
        needles = entry["contains"] if isinstance(entry["contains"], list) else [entry["contains"]]
        return all(
            any(
                (c.get("found") and n in str(c["found"])) or (c["ok"] and c.get("expected") == n)
                for c in matched
            )
            for n in needles
        )
    if "text" in entry:
        return any(c["ok"] and c.get("expected") == entry["text"] for c in matched)
    if "value" in entry:
        return any(c["ok"] and V._numbers_equal(c.get("expected"), entry["value"]) for c in matched)
    if "formula" in entry:
        return any(c["ok"] and c.get("expected") == entry["formula"] for c in matched)
    if "equals" in entry:
        return any(c["ok"] and c.get("expected") == entry["equals"] for c in matched)
    if "count" in entry:
        return any(c["ok"] and c.get("expected") == entry["count"] for c in matched)
    if "structure" in entry:
        return any(c["ok"] and c.get("expected") == entry["structure"] for c in report.checks)
    if "title" in entry:
        ok_title = any(
            c["ok"] and c.get("expected") == entry["title"] and c.get("kind") == "title"
            for c in matched
        )
        needles = entry.get("contains", [])
        ok_contains = all(
            any(c["ok"] and c.get("expected") == n for c in matched) for n in needles
        )
        return ok_title and ok_contains
    raise AssertionError(f"unrecognised truth.json expect entry shape: {entry!r}")


# ------------------------------------------------------- fixtures pass truth.json


@pytest.mark.parametrize(("path", "fmt"), CASES)
def test_fixture_render_validates_ok(path: str, fmt: str) -> None:
    spec = _load_spec(path)
    result = R.render_factory(fmt, spec)
    report = V.validate(spec, fmt, result.data)
    assert report.ok, report.checks


@pytest.mark.parametrize(("path", "fmt"), CASES)
def test_fixture_render_satisfies_every_truth_json_expectation(path: str, fmt: str) -> None:
    spec = _load_spec(path)
    result = R.render_factory(fmt, spec)
    report = V.validate(spec, fmt, result.data)
    expect = TRUTH[_fixture_name(path)]["expect"].get(fmt, [])
    assert expect, f"truth.json has no expect entries for {path}/{fmt}"
    for entry in expect:
        assert _truth_entry_satisfied(report, entry), (entry, report.checks)


def test_truth_json_spoken_numbers_match_the_spec() -> None:
    for path in FIXTURE_SPECS:
        spec = _load_spec(path)
        name = _fixture_name(path)
        assert spec.spoken_numbers == TRUTH[name]["spoken_numbers"]


# ----------------------------------------------------------------- lying renderers


def test_lying_xlsx_dropped_row_is_caught() -> None:
    spec = _load_spec("tests/fixtures/artifacts/specs/butce-tablosu.json")
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Ozet")
    ws.cell(row=1, column=1, value="Kalem")
    ws.cell(row=1, column=2, value="Tutar")
    ws.cell(row=2, column=1, value="Kira")
    ws.cell(row=2, column=2, value=12000)
    # Row 3 (Maaş/45000) is dropped on purpose.
    ws.cell(row=3, column=1, value="Yazılım")
    ws.cell(row=3, column=2, value=8000)
    buf = io.BytesIO()
    wb.save(buf)

    report = V.validate(spec, "xlsx", buf.getvalue())
    assert not report.ok
    assert "sheet:Ozet!A3" in report.failing_refs
    assert "sheet:Ozet!B3" in report.failing_refs


def test_lying_pptx_dropped_slide_is_caught() -> None:
    spec = _load_spec("tests/fixtures/artifacts/specs/q3-sunum.json")
    prs = Presentation()
    layout = prs.slide_layouts[1]
    for slide_spec in spec.slides[:2]:  # the 3rd slide ("Sonuç") is dropped
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = slide_spec.title
        text_frame = slide.placeholders[1].text_frame
        text_frame.clear()
        for i, bullet in enumerate(slide_spec.bullets):
            if i == 0:
                text_frame.text = bullet
            else:
                text_frame.add_paragraph().text = bullet
    buf = io.BytesIO()
    prs.save(buf)

    report = V.validate(spec, "pptx", buf.getvalue())
    assert not report.ok
    assert "s3" in report.failing_refs
    assert "structure:slide_count" in report.failing_refs


def test_lying_docx_corrupted_heading_is_caught() -> None:
    spec = _load_spec("tests/fixtures/artifacts/specs/toplanti-notlari.json")
    document = Document()
    document.add_heading("Giris_WRONG", level=0)  # "Giriş" corrupted
    document.add_paragraph("irrelevant filler paragraph")
    buf = io.BytesIO()
    document.save(buf)

    report = V.validate(spec, "docx", buf.getvalue())
    assert not report.ok
    assert "h1:Giriş" in report.failing_refs


def test_lying_xlsx_cell_that_became_a_real_formula_is_caught_even_when_text_matches() -> None:
    """HIGH security-review finding (ADR-0085 addendum 6): a cell whose stored TEXT
    happens to equal the spec's own expected literal (the spec cell IS the formula
    string) used to pass the plain content-equality check even though the file now
    carries a LIVE formula (``data_type == "f"``), not the literal text asked for. The
    spec itself now refuses to construct such a spec (test_artifact_spec.py); this
    proves the independent validator ALSO catches it, for a renderer that (hypothetic-
    ally, via a bug) still let one through."""
    bad_sheet = Sheet.model_construct(
        name="Ozet",
        columns=["Kalem", "Tutar"],
        rows=[['=HYPERLINK("http://evil","x")', 1]],
        totals=None,
        formulas=None,
    )
    spec = ArtifactSpec.model_construct(
        kind="spreadsheet",
        title="T",
        language="tr",
        sections=None,
        sheets=[bad_sheet],
        slides=None,
        columns=None,
        rows=None,
        style=None,
        spoken_numbers=None,
    )
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Ozet")
    ws.cell(row=1, column=1, value="Kalem")
    ws.cell(row=1, column=2, value="Tutar")
    ws.cell(row=2, column=1, value='=HYPERLINK("http://evil","x")')  # openpyxl: data_type "f"
    ws.cell(row=2, column=2, value=1)
    buf = io.BytesIO()
    wb.save(buf)

    report = V.validate(spec, "xlsx", buf.getvalue())
    assert not report.ok
    assert "sheet:Ozet!A2" in report.failing_refs
    unexpected = [c for c in report.checks if c.get("kind") == "unexpected_formula"]
    assert unexpected and unexpected[0]["ref"] == "sheet:Ozet!A2"
    # The content-equality check on the SAME cell still reports ok=True (the text DOES
    # match) -- the unexpected_formula check is what actually catches this.
    text_checks = [
        c
        for c in report.checks
        if c["ref"] == "sheet:Ozet!A2" and c.get("kind") != "unexpected_formula"
    ]
    assert text_checks and text_checks[0]["ok"] is True


def test_declared_formula_cell_is_not_flagged_as_unexpected() -> None:
    spec = _load_spec("tests/fixtures/artifacts/specs/butce-tablosu.json")
    result = R.render_factory("xlsx", spec)
    report = V.validate(spec, "xlsx", result.data)
    assert report.ok
    assert not any(c.get("kind") == "unexpected_formula" for c in report.checks)


def test_lying_document_missing_paragraph_is_caught_by_ref_text() -> None:
    spec = _load_spec("tests/fixtures/artifacts/specs/hosgeldin-sayfasi.json")
    html = "<!DOCTYPE html><html><body><h1>Hoş geldin</h1></body></html>"
    report = V.validate(spec, "html", html.encode("utf-8"))
    assert not report.ok
    # the section-2 heading and its paragraph text are both missing
    assert any(ref.startswith("h2:") for ref in report.failing_refs)


# ----------------------------------------------------------------- resource bounds


def _zip_bomb_docx(inflated_mb: int = 80) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("word/document.xml", b"0" * (inflated_mb * 1024 * 1024))
    return buf.getvalue()


def test_ooxml_decompression_bomb_is_refused_not_crashed() -> None:
    spec = _load_spec("tests/fixtures/artifacts/specs/toplanti-notlari.json")
    report = V.validate(spec, "docx", _zip_bomb_docx())
    assert not report.ok
    assert report.failing_refs == ["bound:decompression_bound"]


def test_ooxml_too_many_parts_is_refused() -> None:
    spec = _load_spec("tests/fixtures/artifacts/specs/q3-sunum.json")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for i in range(V.MAX_OOXML_PARTS + 1):
            zf.writestr(f"part{i}.xml", "x")
    report = V.validate(spec, "pptx", buf.getvalue())
    assert not report.ok
    assert report.failing_refs == ["bound:too_many_parts"]


def test_oversized_plain_text_input_is_refused() -> None:
    spec = _load_spec("tests/fixtures/artifacts/specs/toplanti-notlari.json")
    huge = b"x" * (V.MAX_INPUT_BYTES + 1)
    report = V.validate(spec, "txt", huge)
    assert not report.ok
    assert report.failing_refs == ["bound:input_too_large"]


def test_oversized_json_input_is_refused_before_parsing() -> None:
    spec = _load_spec("tests/fixtures/artifacts/specs/musteri-listesi.json")
    huge = b"[" + b"1," * (V.MAX_JSON_BYTES) + b"1]"
    report = V.validate(spec, "json", huge)
    assert not report.ok
    assert report.failing_refs == ["bound:too_large"]


def test_pdf_page_bound_is_refused() -> None:
    from pypdf import PdfWriter

    spec = _load_spec("tests/fixtures/artifacts/specs/toplanti-notlari.json")
    writer = PdfWriter()
    for _ in range(V.MAX_PDF_PAGES + 1):
        writer.add_blank_page(width=72, height=72)
    buf = io.BytesIO()
    writer.write(buf)
    report = V.validate(spec, "pdf", buf.getvalue())
    assert not report.ok
    assert report.failing_refs == ["bound:page_bound"]


def _billion_laughs_xlsx() -> bytes:
    """A minimal, well-formed OOXML zip (tiny — never trips the size/ratio bounds
    above) whose ONE worksheet part carries a classic "billion laughs" DOCTYPE. Proves
    LOW security-review finding (ADR-0085 addendum 6): with ``defusedxml`` active,
    openpyxl refuses this outright rather than expanding it — and this module must
    turn that refusal into a failing report, never an unhandled exception."""
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" '
        'ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/worksheets/sheet1.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        "</Types>"
    )
    rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
        'Target="xl/workbook.xml"/></Relationships>'
    )
    wb_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" '
        'Target="worksheets/sheet1.xml"/></Relationships>'
    )
    workbook_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<sheets><sheet name="Sheet1" sheetId="1" r:id="rId1"/></sheets></workbook>'
    )
    sheet_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        "<!DOCTYPE root [\n"
        '<!ENTITY lol "lol">\n'
        '<!ENTITY lol2 "&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;">\n'
        "]>\n"
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<sheetData><row r="1"><c r="A1" t="str"><v>&lol2;</v></c></row></sheetData>'
        "</worksheet>"
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", content_types)
        zf.writestr("_rels/.rels", rels)
        zf.writestr("xl/_rels/workbook.xml.rels", wb_rels)
        zf.writestr("xl/workbook.xml", workbook_xml)
        zf.writestr("xl/worksheets/sheet1.xml", sheet_xml)
    return buf.getvalue()


def test_defusedxml_is_active_in_this_environment() -> None:
    # LOW finding: this was already true TRANSITIVELY (pulled in by fpdf2) before
    # ``defusedxml`` became this project's own explicit runtime dependency; pinning it
    # explicitly is what keeps this assertion true independent of any OTHER package's
    # own future dependency changes.
    import openpyxl.xml.functions as openpyxl_xml

    assert openpyxl_xml.DEFUSEDXML is True


def test_xlsx_entity_expansion_payload_is_refused_not_expanded() -> None:
    spec = _load_spec("tests/fixtures/artifacts/specs/butce-tablosu.json")
    report = V.validate(spec, "xlsx", _billion_laughs_xlsx())
    assert not report.ok
    assert report.failing_refs
    assert report.failing_refs[0].startswith("bound:xml_parse_refused")


def test_not_a_zip_is_refused_not_crashed() -> None:
    spec = _load_spec("tests/fixtures/artifacts/specs/toplanti-notlari.json")
    report = V.validate(spec, "docx", b"this is not a zip file at all")
    assert not report.ok
    assert report.failing_refs == ["bound:not_a_zip"]


def test_malformed_json_is_refused_not_crashed() -> None:
    spec = _load_spec("tests/fixtures/artifacts/specs/musteri-listesi.json")
    report = V.validate(spec, "json", b"{not valid json")
    assert not report.ok
    assert report.failing_refs == ["bound:parse_failed"]


# ---------------------------------------------------------------------- dispatch


def test_validate_refuses_a_format_the_kind_cannot_produce() -> None:
    spec = _load_spec("tests/fixtures/artifacts/specs/butce-tablosu.json")  # spreadsheet
    with pytest.raises(ValueError, match="not valid for kind"):
        V.validate(spec, "pptx", b"irrelevant")
