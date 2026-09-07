"""Generate the M20 document fixtures and their ground truth.

Usage (from services/api, the dev group carries openpyxl + python-pptx):
    uv run python ../../scripts/tests/make-document-fixtures.py ../../services/api/tests/fixtures/documents

Every fixture is deterministic (fixed document properties, zip entries re-written with a
fixed timestamp in sorted order) so a re-run that changes nothing changes no bytes. The
outputs are committed; ``truth.json`` and ``expected/*.extract.json`` are the oracle the
device extractor (C#) and the Cloud Core answer tests are measured against. The reference
scheme (``p3``, ``s4``, ``sheet:Ozet!A5:B5``, ``h2:Kararlar``, ``r7``, ``$.ses``, ``L1-40``)
is the one ADR-0083 fixes: the same string names the place on the device, in Cloud Core's
answer and in the owner's ear.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import sys
import zipfile
from datetime import UTC, datetime
from pathlib import Path

FIXED = datetime(2026, 9, 8, 9, 0, 0, tzinfo=UTC)
ZIP_DATE = (1980, 1, 1, 0, 0, 0)
FONT_DIR = Path(__file__).resolve().parents[2] / "services" / "api" / "app" / "artifacts" / "fonts"


def normalise_zip(path: Path) -> None:
    """Rewrite an OOXML zip with sorted entries and a fixed timestamp (byte-deterministic)."""
    with zipfile.ZipFile(path) as src:
        entries = {name: src.read(name) for name in src.namelist()}
    core = "docProps/core.xml"
    if core in entries:
        # openpyxl stamps dcterms:modified with "now" at save time regardless of wb.properties.
        stamp = FIXED.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
        entries[core] = re.sub(rb'(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:(?:created|modified)>)', rb"\g<1>" + stamp + rb"\g<2>", entries[core])
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for name in sorted(entries):
            info = zipfile.ZipInfo(name, date_time=ZIP_DATE)
            info.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(info, entries[name])
    path.write_bytes(buf.getvalue())


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def norm(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


# --------------------------------------------------------------------------- DOCX
CLAUSES_V1 = [
    ("Madde 1 - Taraflar", "Bu sözleşme, Alp Akbaba (İşveren) ile Örnek Yazılım A.Ş. (Yüklenici) arasında imzalanmıştır."),
    ("Madde 2 - Konu", "Yüklenici, kişisel ajan işletim sisteminin bakım ve geliştirme hizmetlerini üstlenir."),
    ("Madde 3 - Ödeme", "Fatura tarihinden itibaren ödeme süresi 30 gündür."),
    ("Madde 4 - Gizlilik", "Taraflar, sözleşme kapsamında edindikleri bilgileri üçüncü kişilerle paylaşmaz."),
    ("Madde 5 - Süre", "Sözleşme imza tarihinden itibaren bir yıl geçerlidir ve yazılı bildirimle yenilenir."),
]
CLAUSES_V2 = [(h, ("Fatura tarihinden itibaren ödeme süresi 45 gündür." if h.startswith("Madde 3") else t)) for h, t in CLAUSES_V1]


def make_docx(path: Path, clauses: list[tuple[str, str]], version_note: str) -> list[dict]:
    import docx  # python-docx

    doc = docx.Document()
    props = doc.core_properties
    props.created = FIXED.replace(tzinfo=None)
    props.modified = FIXED.replace(tzinfo=None)
    props.author = "PagentOS fixtures"
    props.title = "Hizmet Sözleşmesi"
    blocks: list[dict] = []
    n = 0

    def add(text: str, style: str | None, level: int | None) -> None:
        nonlocal n
        n += 1
        blocks.append({"ref": f"p{n}", "text": text, "kind": "heading" if level else "paragraph", **({"level": level} if level else {})})

    doc.add_heading("Hizmet Sözleşmesi", level=1)
    add("Hizmet Sözleşmesi", "Heading 1", 1)
    doc.add_paragraph(version_note)
    add(version_note, None, None)
    for heading, text in clauses:
        doc.add_heading(heading, level=2)
        add(heading, "Heading 2", 2)
        doc.add_paragraph(text)
        add(text, None, None)
    table = doc.add_table(rows=3, cols=2)
    rows = [("Taraf", "İmza"), ("İşveren", "Alp Akbaba"), ("Yüklenici", "Örnek Yazılım A.Ş.")]
    for r, (a, b) in enumerate(rows):
        table.cell(r, 0).text = a
        table.cell(r, 1).text = b
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    normalise_zip(path)
    return blocks + [{"ref": "t1", "kind": "table", "text": "\n".join("\t".join(r) for r in rows), "rows": [list(r) for r in rows]}]


# --------------------------------------------------------------------------- XLSX
def make_xlsx(path: Path) -> tuple[list[dict], dict]:
    from openpyxl import Workbook

    wb = Workbook()
    wb.properties.created = FIXED.replace(tzinfo=None)
    wb.properties.modified = FIXED.replace(tzinfo=None)
    wb.properties.creator = "PagentOS fixtures"
    ws = wb.active
    ws.title = "Ozet"
    ozet = [("Kalem", "Tutar"), ("Kira", 12000), ("Maaş", 45000), ("Yazılım", 8000), ("Toplam", 65000)]
    for row in ozet:
        ws.append(row)
    ws["A6"] = "KDV"
    ws["B6"] = "=B5*0.2"
    detay = wb.create_sheet("Detay")
    months = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]
    detay.append(("Ay", "Gider"))
    for i, m in enumerate(months):
        detay.append((m, 5000 + i * 100))
    wb.save(path)
    normalise_zip(path)

    blocks: list[dict] = []
    for r, row in enumerate(ozet, start=1):
        blocks.append({"ref": f"sheet:Ozet!A{r}:B{r}", "kind": "row", "sheet": "Ozet", "text": "\t".join(str(v) for v in row)})
    blocks.append({"ref": "sheet:Ozet!A6:B6", "kind": "row", "sheet": "Ozet", "text": "KDV\t=B5*0.2", "formulas": {"B6": "B5*0.2"}})
    blocks.append({"ref": "sheet:Detay!A1:B1", "kind": "row", "sheet": "Detay", "text": "Ay\tGider"})
    for i, m in enumerate(months, start=2):
        blocks.append({"ref": f"sheet:Detay!A{i}:B{i}", "kind": "row", "sheet": "Detay", "text": f"{m}\t{5000 + (i - 2) * 100}"})
    structure = {"sheets": [{"name": "Ozet", "rows": 6, "columns": 2}, {"name": "Detay", "rows": 13, "columns": 2}], "formulas": [{"ref": "sheet:Ozet!B6", "formula": "B5*0.2"}]}
    return blocks, structure


# --------------------------------------------------------------------------- PPTX
SLIDES = [
    ("Q3 Özeti", ["Gelir hedefi aşıldı", "Üç yeni müşteri", "Bir gecikmiş teslimat"]),
    ("Gelirler", ["Toplam gelir 1,2 milyon TL", "Abonelik payı yüzde 60"]),
    ("Giderler", ["Personel gideri sabit", "Bulut maliyeti yüzde 8 arttı"]),
    ("Müşteri Kazanımı", ["Üç kurumsal müşteri", "Dönüşüm oranı yüzde 12"]),
    ("Riskler", ["Tedarikçi bağımlılığı", "Kur dalgalanması"]),
    ("Hedefler", ["Q4'te iki ürün lansmanı", "Destek süresini yarıya indirmek"]),
    ("Teşekkürler", ["Sorular?"]),
]


def make_pptx(path: Path) -> tuple[list[dict], dict]:
    from pptx import Presentation

    prs = Presentation()
    prs.core_properties.created = FIXED.replace(tzinfo=None)
    prs.core_properties.modified = FIXED.replace(tzinfo=None)
    prs.core_properties.author = "PagentOS fixtures"
    prs.core_properties.title = "Q3 Sunumu"
    layout = prs.slide_layouts[1]  # title + content
    for title, bullets in SLIDES:
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = title
        body = slide.placeholders[1].text_frame
        body.text = bullets[0]
        for b in bullets[1:]:
            body.add_paragraph().text = b
    prs.save(path)
    normalise_zip(path)
    blocks = [{"ref": f"s{i}", "kind": "slide", "title": t, "text": "\n".join([t, *b])} for i, (t, b) in enumerate(SLIDES, start=1)]
    return blocks, {"slides": [{"index": i, "title": t} for i, (t, _) in enumerate(SLIDES, start=1)], "slide_count": len(SLIDES)}


# --------------------------------------------------------------------------- PDF
PAGES = [
    ["Yıllık Rapor 2026", "Bu rapor kişisel ajan işletim sisteminin yıllık durumunu özetler."],
    ["Bölüm 1 - Altyapı", "Bulut çekirdeği Hetzner üzerinde mavi-yeşil dağıtımla çalışır."],
    ["Bölüm 2 - Belgeler", "Üçüncü sayfada yer alan bu cümle referans testidir.", "Belge zekâsı her yanıtta dosyayı ve yeri adıyla söyler."],
    ["Bölüm 3 - Ses", "Türkçe ses yolu ilk sınıf vatandaştır."],
    ["Sonuç", "Sistem sahibinin bakımcı olmasını gerektirmez."],
]


def make_pdf(path: Path, font_dir: Path) -> list[dict]:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_creation_date(FIXED)
    pdf.set_creator("PagentOS fixtures")
    pdf.set_title("Yıllık Rapor 2026")
    pdf.add_font("DejaVu", "", str(font_dir / "DejaVuSans.ttf"))
    pdf.add_font("DejaVu", "B", str(font_dir / "DejaVuSans-Bold.ttf"))
    for lines in PAGES:
        pdf.add_page()
        pdf.set_font("DejaVu", "B", 16)
        pdf.cell(0, 12, lines[0], new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("DejaVu", "", 12)
        for line in lines[1:]:
            pdf.multi_cell(0, 8, line)
            pdf.ln(2)
    pdf.output(str(path))
    return [{"ref": f"p{i}", "kind": "page", "match": "contains", "text": " ".join(lines), "contains": lines} for i, lines in enumerate(PAGES, start=1)]


# --------------------------------------------------------------------------- text-like
NOTLAR = """# Toplantı Notları

Tarih: 8 Eylül 2026. Katılımcılar: Alp, sistem.

## Giriş

M20 belge zekâsı kapsamı konuşuldu. Dosyalar sahibin makinesinde kalır.

## Kararlar

- Bütçe onaylandı.
- Her yanıt dosyayı ve yeri adıyla söyler.
- Silme yeteneği bu aşamada yoktur.

## Sonraki adımlar

Sabitler üretilir, cihaz çıkarımı doğrulanır, ses derlemi genişletilir.
"""

VERI_ROWS = [("Ad", "Şehir", "Yaş", "Tutar", "Durum")] + [
    ("Ayşe", "Ankara", 34, 1200, "aktif"),
    ("Mehmet", "İstanbul", 41, 980, "pasif"),
    ("Zeynep", "Bursa", 29, 1500, "aktif"),
    ("Can", "Antalya", 52, 700, "aktif"),
    ("Elif", "Adana", 37, 1100, "pasif"),
    ("Kerem", "İzmir", 45, 1320, "aktif"),
    ("Deniz", "Trabzon", 31, 890, "aktif"),
    ("Selin", "Eskişehir", 27, 1010, "pasif"),
    ("Burak", "Konya", 39, 1230, "aktif"),
    ("Naz", "Samsun", 33, 760, "aktif"),
]

AYARLAR = {
    "surum": "0.2.0",
    "ses": {"dil": "tr-TR", "hiz": 1.0, "ses_adi": "alloy"},
    "cihaz": {"ad": "MAIL", "operator": True},
    "kokler": ["Documents", "Desktop", "Downloads"],
}

KOD = '''"""Örnek modül: bütçe hesapları."""


def hesapla(kira: int, maas: int, yazilim: int) -> int:
    """Üç kalemin toplamı."""
    return kira + maas + yazilim


def kdv(tutar: int, oran: float = 0.2) -> float:
    return tutar * oran


class Butce:
    def __init__(self, kalemler: dict[str, int]) -> None:
        self.kalemler = kalemler

    def toplam(self) -> int:
        return sum(self.kalemler.values())
'''


def md_blocks(text: str) -> list[dict]:
    blocks: list[dict] = []
    current: dict | None = None
    body: list[str] = []

    def flush() -> None:
        if current is not None:
            current["text"] = norm("\n".join(body))
            blocks.append(current)

    for line in text.splitlines():
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            flush()
            level = len(m.group(1))
            current = {"ref": f"h{level}:{m.group(2).strip()}", "kind": "section", "level": level, "title": m.group(2).strip()}
            body = []
        else:
            body.append(line)
    flush()
    return blocks


def csv_blocks(rows: list[tuple]) -> tuple[list[dict], dict]:
    lines = []
    for row in rows:
        lines.append(",".join(str(v) for v in row))
    blocks = [{"ref": f"r{i}", "kind": "row", "text": line} for i, line in enumerate(lines, start=1)]
    return blocks, {"columns": list(rows[0]), "rows": len(rows) - 1, "delimiter": ","}


def json_blocks(obj: dict) -> tuple[list[dict], dict]:
    blocks = [{"ref": f"$.{k}", "kind": "key", "text": json.dumps(v, ensure_ascii=False, separators=(",", ":"))} for k, v in obj.items()]
    return blocks, {"keys": list(obj.keys())}


def source_blocks(text: str) -> tuple[list[dict], dict]:
    lines = text.splitlines()
    blocks = []
    for start in range(0, len(lines), 40):
        chunk = lines[start:start + 40]
        blocks.append({"ref": f"L{start + 1}-{start + len(chunk)}", "kind": "lines", "text": "\n".join(chunk)})
    symbols = []
    for i, line in enumerate(lines, start=1):
        m = re.match(r"^\s*(?:def|class)\s+([A-Za-z_][A-Za-z0-9_]*)", line)
        if m:
            symbols.append({"name": m.group(1), "line": i})
    return blocks, {"symbols": symbols, "lines": len(lines), "language": "python"}


# --------------------------------------------------------------------------- main
def main(out: Path) -> None:
    out.mkdir(parents=True, exist_ok=True)
    expected_dir = out / "expected"
    expected_dir.mkdir(exist_ok=True)
    files: list[dict] = []

    def record(rel: str, kind: str, blocks: list[dict], structure: dict | None = None, *, title: str | None = None, match: str = "equals_normalised") -> None:
        p = out / rel
        entry = {"path": rel, "kind": kind, "size": p.stat().st_size, "sha256": sha256(p), "blocks": len(blocks)}
        if title:
            entry["title"] = title
        files.append(entry)
        (expected_dir / (rel.replace("/", "__") + ".extract.json")).write_text(
            json.dumps({"path": rel, "kind": kind, "match": match, "title": title, "structure": structure or {}, "blocks": blocks}, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8", newline="\n",
        )

    b = make_docx(out / "sozlesmeler/2025/sozlesme.docx", CLAUSES_V1, "Sürüm 1 (2025).")
    record("sozlesmeler/2025/sozlesme.docx", "docx", b, {"headings": [x["text"] for x in b if x["kind"] == "heading"], "tables": 1}, title="Hizmet Sözleşmesi")
    b = make_docx(out / "sozlesmeler/2026/sozlesme.docx", CLAUSES_V2, "Sürüm 2 (2026).")
    record("sozlesmeler/2026/sozlesme.docx", "docx", b, {"headings": [x["text"] for x in b if x["kind"] == "heading"], "tables": 1}, title="Hizmet Sözleşmesi")
    b, s = make_xlsx(out / "butce-2026.xlsx")
    record("butce-2026.xlsx", "xlsx", b, s, title="butce-2026")
    b, s = make_pptx(out / "sunum-q3.pptx")
    record("sunum-q3.pptx", "pptx", b, s, title="Q3 Sunumu")
    b = make_pdf(out / "rapor.pdf", FONT_DIR)
    record("rapor.pdf", "pdf", b, {"pages": len(PAGES)}, title="Yıllık Rapor 2026", match="contains")
    (out / "notlar.md").write_text(NOTLAR, encoding="utf-8", newline="\n")
    b = md_blocks(NOTLAR)
    record("notlar.md", "md", b, {"headings": [x["title"] for x in b]}, title="Toplantı Notları")
    with (out / "veri.csv").open("w", encoding="utf-8", newline="\n") as fh:
        for row in VERI_ROWS:
            fh.write(",".join(str(v) for v in row) + "\n")
    b, s = csv_blocks(VERI_ROWS)
    record("veri.csv", "csv", b, s)
    (out / "ayarlar.json").write_text(json.dumps(AYARLAR, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
    b, s = json_blocks(AYARLAR)
    record("ayarlar.json", "json", b, s)
    (out / "kod.py").write_text(KOD, encoding="utf-8", newline="\n")
    b, s = source_blocks(KOD)
    record("kod.py", "source", b, s)

    truth = {
        "generated_by": "scripts/tests/make-document-fixtures.py",
        "reference_scheme": {
            "docx": "p<n> per non-empty paragraph in order (headings included, kind=heading with level); t<n> per table",
            "xlsx": "sheet:<name>!A<r>:<lastcol><r> per non-empty row; formulas reported as text with a leading '='",
            "pptx": "s<n> per slide; text = title then body paragraphs",
            "pdf": "p<n> per page; match=contains (layout engines differ in whitespace)",
            "md": "h<level>:<title> per heading section; text = the section body",
            "csv": "r<n> per line including the header line",
            "json": "$.<key> per top-level key; text = compact JSON of the value",
            "source": "L<a>-<b> per 40-line chunk; structure.symbols from def/class lines",
        },
        "files": files,
        "questions": [
            {"id": "pdf-page-3", "file": "rapor.pdf", "question": "Üçüncü sayfada ne yazıyor?", "expect_ref": "p3", "expect_contains": "referans testidir"},
            {"id": "xlsx-total", "file": "butce-2026.xlsx", "question": "Toplam bütçe ne kadar?", "expect_ref": "sheet:Ozet!A5:B5", "expect_contains": "65000"},
            {"id": "xlsx-formula", "file": "butce-2026.xlsx", "question": "KDV nasıl hesaplanıyor?", "expect_ref": "sheet:Ozet!A6:B6", "expect_contains": "B5*0.2"},
            {"id": "pptx-slide-4", "file": "sunum-q3.pptx", "question": "Dördüncü slaytın başlığı ne?", "expect_ref": "s4", "expect_contains": "Müşteri Kazanımı"},
            {"id": "pptx-count", "file": "sunum-q3.pptx", "question": "Sunumda kaç slayt var?", "expect_structure": {"slide_count": 7}},
            {"id": "docx-clause-3", "file": "sozlesmeler/2026/sozlesme.docx", "question": "Ödeme süresi kaç gün?", "expect_ref": "p8", "expect_contains": "45 gün"},
            {"id": "md-decisions", "file": "notlar.md", "question": "Kararlar bölümünde ne var?", "expect_ref": "h2:Kararlar", "expect_contains": "Bütçe onaylandı"},
            {"id": "csv-row-7", "file": "veri.csv", "question": "Kerem hangi şehirde?", "expect_ref": "r7", "expect_contains": "İzmir"},
            {"id": "json-dil", "file": "ayarlar.json", "question": "Ses dili ne?", "expect_ref": "$.ses", "expect_contains": "tr-TR"},
            {"id": "py-symbol", "file": "kod.py", "question": "hesapla fonksiyonu kaçıncı satırda?", "expect_structure": {"symbol": "hesapla", "line": 4}},
        ],
        "comparisons": [
            {
                "id": "sozlesme-v1-v2",
                "a": "sozlesmeler/2025/sozlesme.docx",
                "b": "sozlesmeler/2026/sozlesme.docx",
                "same_title": "Hizmet Sözleşmesi",
                "changed_refs": ["p2", "p8"],
                "changed_clause": {"ref": "p8", "a_contains": "30 gün", "b_contains": "45 gün"},
                "unchanged_refs": ["p1", "p3", "p4", "p5", "p6", "p7", "p9", "p10", "p11", "p12", "t1"],
            }
        ],
        "common_points": [
            {"id": "budget-across", "files": ["butce-2026.xlsx", "kod.py", "notlar.md"], "expect_terms": ["bütçe"]},
        ],
    }
    (out / "truth.json").write_text(json.dumps(truth, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
    print(f"wrote {len(files)} fixtures to {out}")


if __name__ == "__main__":
    main(Path(sys.argv[1]).resolve())
